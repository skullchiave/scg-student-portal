#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""import_fmt_xlsx.py — 「課題登録FMT」の Excel を、学生ポータルの形へ変換する（2026-09-10）

これは **作問の入口側** を読む取り込み口。ヨリソルの書き出しCSVを読む
`import_yorisol.py`（出口側）とは別物なので、混ぜないこと。

  入口 = 先生が問題を書く Excel（このスクリプトが読む）
  出口 = ヨリソルから書き出した CSV（import_yorisol.py が読む）

**DB へは接続しない。** 出力はファイル（JSON と SQL）で、流し込むかどうかは人が決める。

■ 「課題登録FMT」とは
  ドライブの教材フォルダにある作成用 Excel の 1 枚目
  「250507から_課題登録FMT_コピーして使用」がこの形式。先生はこのシートをコピーして
  1回ぶんのテストを書く。つまり **1シート = 1回のテスト = quiz_sets 1件**。
  ★2026-09-10 に判明: ファイル名に「JapanGo_」と付いていても中身はこの FMT で、
    ヨリソル形式/JapanGo形式という2つの形は**存在しなかった**。読む形はこれ1つでよい。

■ 列（14列・位置で読む）
   1 問題番号 / 2 問題文1 / 3 問題文2 / 4 添付ファイル名 / 5-9 選択肢1〜5 /
   10 解説 / 11 解答 / 12 カテゴリ / 13 配点 / 14 (空)
  ★**見出しの文字列に頼らず、列の位置で読む。**
    実データでは「④10-12」シートだけ「解答」の見出しが消え、「カテゴリ」が
    「文法読解」に書き換わっていた。見出しで読むと、そのシートだけ黙って壊れる。
    見出しは検査だけして、違っていたら警告に出す。

■ 拾わないもの
  ・問題番号が数値でない行（実データでは各シート22行目の「集計」）
  ・シート名に FMT を含むシート（テンプレート本体。見本の10問が入っている）

■ 画像について（2026-09-10 きあ判断）
  設問には「添付ファイル名」列があり、実データでも 200問中2問で使われている
  （7-9-⑮ゴミ出し.png など）。**今のポータルは画像を出さない**。
  ただし **データは捨てない**＝JSON には image_name を必ず書き出す。
  ★画像そのものは Supabase に入れない。持つのはファイル名だけで、出すときは
    学校の Google ドライブ側の配信口を使う（顔写真を入れないと決めたのと同じ考え方）。

■ 足りなかった4つ（2026-09-10 に列を足した）
  画像・カテゴリ・配点・解説は、以前は JSON にだけ残して SQL には出していなかった
  ＝**DBに入れた時点で消えていた**。db/2026-09-10_question_columns.sql で列を足したので、
  既定（--schema 2026-09-10）ではこの4つも SQL に書く。
  🔴 **解説だけ入れ先が違う**＝questions ではなく question_answers（教師のみ読める表）。
     questions は「公開中の回なら学生が読める」ので、置くと受験前に解説＝正解が漏れる。
  列を足していない相手へ書き出すときは --schema 2026-09-06（★その4つは捨てられる）。

■ 使い方
  py -X utf8 scripts\import_fmt_xlsx.py --xlsx "...\★ヨリソル_まとめテストⅠ（作成用）.xlsx"
      … 全シートを読んで、結果だけ表示（ファイルは書かない）

  py -X utf8 scripts\import_fmt_xlsx.py --xlsx "...(作成用).xlsx" ^
      --sheet "①1-3 （新）" --title-prefix "つなぐ日本語初級 まとめテスト" ^
      --out-json tmp\q.json --out-sql tmp\q.sql

■ 出力（ポータルのスキーマ・既定 --schema 2026-09-10）
  quiz_sets(title, lesson, is_open)                    … 1シートにつき1件。**is_open は false**
  questions(quiz_set_id, seq, prompt,
            image_name, category, points)              … 問題文1と問題文2を改行でつなぐ
  question_choices(question_id, idx, label)            … 空でない選択肢を1から詰める
  question_answers(question_id, <正解列>, explanation) … 値は「何番目が正解か」（1始まり）
  ⓘ quiz_set_questions（②と①をつなぐ表）への登録は書かない。
    db/2026-09-10_four_layers.sql のトリガ questions_sync_set_link が自動で入れる。
"""
from __future__ import annotations

import argparse
import collections
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    print("openpyxl が要ります:  py -m pip install openpyxl", file=sys.stderr)
    raise

# ---- 列の位置（1始まり）。★見出しではなくここで読む -------------------------
COL_NO, COL_TEXT1, COL_TEXT2, COL_IMAGE = 1, 2, 3, 4
COL_CHOICE_FIRST, N_CHOICE_COLS = 5, 5          # 選択肢1〜5 → 5,6,7,8,9
COL_EXPLAIN, COL_ANSWER, COL_CATEGORY, COL_POINTS = 10, 11, 12, 13

EXPECTED_HEADERS = {
    COL_NO: "問題番号", COL_TEXT1: "問題文1", COL_TEXT2: "問題文2", COL_IMAGE: "添付ファイル名",
    5: "選択肢1", 6: "選択肢2", 7: "選択肢3", 8: "選択肢4", 9: "選択肢5",
    COL_EXPLAIN: "解説", COL_ANSWER: "解答", COL_CATEGORY: "カテゴリ", COL_POINTS: "配点",
}

# question_answers の正解列名。ライブのスキーマと違ったらここだけ直す（--answer-col でも指定可）
DEFAULT_ANSWER_COL = "correct_idx"
# 出力するスキーマの世代。★2026-09-10 に questions へ image_name / category / points、
# question_answers へ explanation を足した（db/2026-09-10_question_columns.sql）。
# その SQL をまだ流していない相手へ書き出すときだけ 2026-09-06 を指定する。
SCHEMA_NEW, SCHEMA_OLD = "2026-09-10", "2026-09-06"
# 1設問あたりの選択肢の上限。question_choices の check (idx between 1 and 12) と合わせること
MAX_CHOICES = 12
# テンプレート本体とみなすシート名の目印
TEMPLATE_MARK = "FMT"
# シート名の頭に付く丸数字（順序を表すだけなので、タイトルからは外す）
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"

# ★ルビの書き方（2026-09-10 に実データで判明）: `${昨日}(きのう)` のように書かれている。
#   239本のブックに「ルビあり」「ルビなし」の2枚があり、**きあ判断でルビあり版を採用**。
#   ⚠ ルビを外してもルビなし版と完全には一致しない（一致率82.4%）。
#     漢字をひらがなに開いた箇所があり、しかも**両方の版がそれぞれ別の誤字を持っている**。
#     ＝ルビあり版を正と決めることで、表記が1本にそろう。
RUBY = re.compile(r"\$\{([^}]*)\}\(([^)]*)\)")


def apply_ruby(s: str, mode: str) -> str:
    """keep=記法のまま / strip=ルビを消す / html=<ruby> タグにする"""
    if not s or mode == "keep":
        return s
    if mode == "strip":
        return RUBY.sub(r"\1", s)
    if mode == "html":
        return RUBY.sub(r"<ruby>\1<rt>\2</rt></ruby>", s)
    raise ValueError(mode)


# ---------------------------------------------------------------- 小道具
def cell_text(ws, row: int, col: int) -> str:
    v = ws.cell(row=row, column=col).value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def as_int(v) -> int | None:
    """数値なら int にする。文字の '3' も数値として受ける（'集計' は None）。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else None
    s = str(v).strip().translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    return int(s) if re.fullmatch(r"\d+", s) else None


def lesson_of(sheet_name: str) -> str:
    """シート名から「何課ぶんか」を取り出す（'①1-3 （新）' → '1-3'）。

    ★範囲が取れない形（チェックテストの '1-①' や 'オリエン'）は**シート名をそのまま**返す。
      lesson は学生画面のバッジに出る値で、空にすると「Quiz」としか出ず、
      どの回を開いているのか分からなくなるため。
    """
    m = re.search(r"(\d+)\s*[-‐−–~〜]\s*(\d+)", sheet_name)
    if m:
        return f"{int(m.group(1))}-{int(m.group(2))}"
    return sheet_name.strip()


def title_of(sheet_name: str, prefix: str) -> str:
    body = sheet_name.lstrip(CIRCLED).strip()
    return f"{prefix} {body}".strip() if prefix else (body or sheet_name)


# ---------------------------------------------------------------- どの列から始まるか
def anchor_offset(ws) -> int:
    """FMT が**何列目から始まるか**を返す（0 = 1列目から＝素のFMT）。

    ★実データには **FMT の前に1列増えている版**がある（2026-09-10 に331本を調べて判明）:
      ・`課番号`＋FMT … スピードマスター／JLPT対策／完全マスター／まとめドリルの「ルビあり」等
      ・`設問`  ＋FMT … 文法チェックテストⅠ の「全体」シート
    位置で読む方針は変えない（見出しは現場で壊れているため）。
    **どこから始まるかだけ**を「問題番号」の見出しで決める。見つからなければ1列目から。
    """
    for j in range(1, 6):
        v = ws.cell(row=1, column=j).value
        if v is not None and str(v).strip() == "問題番号":
            return j - COL_NO
    return 0


# ---------------------------------------------------------------- 計算結果が入っているか
def stale_cache(wsv, wsf, sheet_name: str, off: int = 0) -> str | None:
    """🔴 **数式はあるのに「計算結果」が入っていない**ブックを名指しする。

    Excel が保存したブックは、数式のセルに計算結果も一緒に持っている。
    ところが **openpyxl で開いて保存し直すと、その計算結果だけが消える**（数式は残る）。
    こちらは計算結果のほうを読む（`data_only=True`）ので、そうなったブックは
    **問題番号が全部空に見えて、1問も取り込めないまま静かに終わる**。

    ★2026-09-10 に実際にやらかした：1セル足すために openpyxl で保存し直したら、
      99問が10問になった。しかも「数式そのもの」を突き合わせる検査をしていたので
      「1セルだけ変わった」と出て、壊れたことに気づけなかった。
      **同じことは作問側が別のツールで保存した場合にも起きる。**
    """
    n_formula = n_missing = 0
    for i in range(2, wsf.max_row + 1):
        f = wsf.cell(row=i, column=COL_NO + off).value
        if isinstance(f, str) and f.startswith("="):
            n_formula += 1
            if wsv.cell(row=i, column=COL_NO + off).value in (None, ""):
                n_missing += 1
    if n_formula and n_missing == n_formula:
        return (f"🔴 [{sheet_name}] 問題番号が数式（{n_formula}行）なのに、計算結果が入っていない"
                f"＝このままでは1問も取り込めない。**Excel で開いて保存し直してから**やり直すこと"
                f"（openpyxl などで書き換えたブックはこうなる）")
    return None


# ---------------------------------------------------------------- 変換（1シート）
def build_sheet(ws, sheet_name: str, prefix: str, ruby: str = "keep") -> tuple[dict, list[str]]:
    warn: list[str] = []
    tag = f"[{sheet_name}]"
    off = anchor_offset(ws)            # ★FMT が何列目から始まるか
    if off:
        warn.append(f"{tag} FMT が {off + 1} 列目から始まっている"
                    f"（1列目は「{cell_text(ws, 1, 1)}」）＝その分ずらして読む")

    # 見出しの検査。★読むのは位置なので、ここは「気づくため」だけに使う
    hit = 0
    for col, want in EXPECTED_HEADERS.items():
        got = cell_text(ws, 1, col + off)
        if got == want:
            hit += 1
        else:
            warn.append(f"{tag} {col + off}列目の見出しが「{got or '(空)'}」＝想定は「{want}」"
                        f"（位置で読むので取り込みは続けるが、列がずれていないか目で確かめること）")

    # ★そもそも FMT でないシートは触らない。
    #   実データの「どんどんつながる漢字練習帳 中級編 7〜13回目一覧」がこれで、
    #   見出しは 回（課）｜問題番号｜問題文1｜問題文2｜選択肢1〜4｜解答 の**10列版**だった
    #   （添付ファイル名の列が無い＝以降が1列ずつずれる）。しかも中身は問題文だけの一覧表。
    #   位置で読むと選択肢を拾えず「選択肢が0個」が 140 件並ぶ＝本物の不備が埋もれる。
    if hit < 9:
        return {"sheet": sheet_name, "title": title_of(sheet_name, prefix),
                "lesson": lesson_of(sheet_name), "is_open": False,
                "questions": [], "dropped": [],
                "error": f"課題登録FMT ではない（見出しの一致 {hit}/{len(EXPECTED_HEADERS)}）"}, \
               [f"{tag} 課題登録FMT ではないので触らない（見出しの一致 {hit}/{len(EXPECTED_HEADERS)}）"]

    items: list[dict] = []
    dropped: list[dict] = []          # ★取り込めなかった問題。**JSON にも残す**
    skipped_rows: list[str] = []
    unwritten: list[int] = []         # 問題番号だけあって中身が無い行＝テンプレートの使い残し

    for row in range(2, ws.max_row + 1):
        raw_no = ws.cell(row=row, column=COL_NO + off).value
        if raw_no is None or str(raw_no).strip() == "":
            continue                      # 空行は静かに飛ばす
        seq = as_int(raw_no)
        if seq is None:
            # ★黙って落とさない。実データでは各シート22行目の「集計」がここに来る
            skipped_rows.append(f"{row}行目「{str(raw_no).strip()[:20]}」")
            continue

        # ★問題番号（と配点）だけで、中身が1つも無い行＝**テンプレートの使い残し**。
        #   実データでは 929 行がこれだった（例: スピードマスターは20問の枠に7問だけ書いてある）。
        #   これを「不備」に混ぜると本物の不備が埋もれるので、**別に数える**。
        if not any(cell_text(ws, row, c + off) for c in
                   (COL_TEXT1, COL_TEXT2, COL_IMAGE, COL_EXPLAIN, COL_ANSWER, COL_CATEGORY)) \
                and not any(cell_text(ws, row, COL_CHOICE_FIRST + off + k) for k in range(N_CHOICE_COLS)):
            unwritten.append(seq)
            continue

        t1 = apply_ruby(cell_text(ws, row, COL_TEXT1 + off), ruby)
        t2 = apply_ruby(cell_text(ws, row, COL_TEXT2 + off), ruby)
        prompt = "\n".join(x for x in (t1, t2) if x)
        image = cell_text(ws, row, COL_IMAGE + off)
        explain = cell_text(ws, row, COL_EXPLAIN + off)
        category = cell_text(ws, row, COL_CATEGORY + off)
        points = as_int(ws.cell(row=row, column=COL_POINTS + off).value)

        # 選択肢: 空でないものを 1 から詰める。★途中が空なら、ずれの疑いとして報告する
        def drop(reason: str, _seq=seq, _row=row) -> None:
            """★落としたことを2か所に残す＝実行時の警告と、JSON の dropped。
            警告だけだと画面を流れて消えるので、**後から JSON を見た人には見えなくなる**。
            とくに**最後の問題が落ちた場合は番号の抜けにならない**ので、
            出来上がりだけ眺めても気づけない（⑩28-30 の問20 が実際にこれだった）。"""
            warn.append(f"{tag} 問{_seq}: {reason}")
            dropped.append({"seq": _seq, "row": _row, "reason": reason})

        raw = [apply_ruby(cell_text(ws, row, COL_CHOICE_FIRST + off + k), ruby)
               for k in range(N_CHOICE_COLS)]
        labels = [s for s in raw if s]
        first_blank = next((k for k, s in enumerate(raw) if not s), N_CHOICE_COLS)
        if any(raw[k] for k in range(first_blank + 1, N_CHOICE_COLS)):
            drop(f"選択肢が飛んでいる（{[bool(s) for s in raw]}）＝正解の番号とずれる恐れがある")
            continue

        if not prompt:
            drop("問題文1も問題文2も空")
            continue
        if len(labels) < 2:
            drop(f"選択肢が {len(labels)} 個（2個以上ないと出題できない）")
            continue
        if len(labels) > MAX_CHOICES:
            drop(f"選択肢が {len(labels)} 個（上限 {MAX_CHOICES} 個）")
            continue

        ans = as_int(ws.cell(row=row, column=COL_ANSWER + off).value)
        if ans is None:
            drop("解答が空か、数字でない")
            continue
        if not (1 <= ans <= len(labels)):
            drop(f"解答が {ans} だが選択肢は {len(labels)} 個しかない")
            continue

        # ★同じ文言の選択肢が並んでいたら、問題として成立していない（実データの ⑩28-30 問3）。
        #   ただし落とすかどうかは**正解がその中にあるか**で変わる。
        #   ・正解が重複の外 → 出題はできる（正解は選べる）ので、取り込んで警告だけ
        #   ・正解が重複の中 → 同じものを選んでも片方しか正解にならない＝**不公平なので落とす**
        dups = sorted(lab for lab, c in collections.Counter(labels).items() if c > 1)
        if dups:
            if labels[ans - 1] in dups:
                drop(f"選択肢が重複していて、しかも正解がその中にある（{dups}）")
                continue
            warn.append(f"{tag} 問{seq}: 選択肢が重複している（{dups}）"
                        f"＝問題として成立していない。作問側に直してもらうこと")

        if image:
            # ★落とさない。今の画面は出さないだけで、名前は JSON に残す
            warn.append(f"{tag} 問{seq}: 画像つきの設問（{image}）"
                        f"＝いまの画面は画像を出さない。文字だけで意味が通るか確認すること")

        items.append({
            "seq": seq,
            "prompt": prompt,
            "choices": labels,
            "correct_idx": ans,
            "image_name": image or None,
            "category": category or None,
            "points": points,
            "explanation": explain or None,
        })

    if skipped_rows:
        warn.append(f"{tag} 問題番号が数字でないので飛ばした行: {'、'.join(skipped_rows)}"
                    f"（「集計」なら想定どおり）")
    if unwritten:
        warn.append(f"{tag} まだ書かれていない問題が {len(unwritten)} 問（問"
                    f"{unwritten[0]}〜{unwritten[-1]}）＝問題番号の枠だけがある状態。"
                    f"**作問の不備ではなく、作りかけ**")

    seqs = [i["seq"] for i in items]
    dup = [s for s, n in collections.Counter(seqs).items() if n > 1]
    if dup:
        warn.append(f"{tag} 問題番号が重複: {sorted(dup)}")

    # ★同じテストの中で配点がそろっていない（2026-09-11 追加）
    #   実測（2026-09-11・教材331本）: 設問のある667シートのうち **662枚（99.3%）は全問おなじ配点**。
    #   本当に2種類以上あるのは**5枚だけ**。★空欄は「違う配点」ではなく「書かれていない」なので数に入れない
    #   （入れると13枚ぶん増えて18枚になり、見てほしい5枚が埋もれる）。
    #   **採点には使っていない**ので取り込みは止めず、作問側に見てもらうために出すだけ。
    #   意図的な傾斜配点ならそのままでよい。
    pts = [i["points"] for i in items if i["points"] is not None]
    if len(set(pts)) > 1:
        dist = collections.Counter(pts)
        rare = min(dist, key=lambda v: (dist[v], v))     # いちばん少ない配点＝打ち間違いの候補
        detail = "、".join(f"{v}点が{n}問" for v, n in sorted(dist.items()))
        warn.append(f"{tag} 同じテストの中で配点がそろっていない（{detail}）"
                    f"＝{rare}点の問だけ違います。打ち間違いでなければそのままで構いません")

    return {
        "sheet": sheet_name,
        "title": title_of(sheet_name, prefix),
        "lesson": lesson_of(sheet_name),
        "is_open": False,
        "questions": items,
        "dropped": dropped,          # ★取り込めなかった問題（空なら [] を書く。無かった証拠になる）
        "unwritten": unwritten,      # まだ書かれていない問題番号＝作りかけ（不備とは別に数える）
    }, warn


# ---------------------------------------------------------------- 変換（ブック全体）
def build(path: Path, only: list[str] | None, exclude: list[str] | None,
          prefix: str, include_template: bool, ruby: str = "keep") -> tuple[list[dict], list[str]]:
    wb = load_workbook(path, data_only=True)
    wbf = load_workbook(path, data_only=False)   # ★数式そのもの。計算結果の有無を見るためだけに使う
    warn: list[str] = []
    sets: list[dict] = []

    for name in wb.sheetnames:
        if only and name not in only:
            continue
        if exclude and name in exclude:
            continue
        if not include_template and TEMPLATE_MARK in name:
            continue                      # テンプレート本体（見本10問）は取り込まない
        stale = stale_cache(wb[name], wbf[name], name, anchor_offset(wb[name]))
        if stale:
            # ★黙って0問にしない。原因を名指しして、そのシートは触らない
            warn.append(stale)
            sets.append({"sheet": name, "title": title_of(name, prefix), "lesson": lesson_of(name),
                         "is_open": False, "questions": [], "dropped": [],
                         "error": "計算結果が入っていない"})
            continue
        s, w = build_sheet(wb[name], name, prefix, ruby)
        warn += w
        if not s["questions"]:
            warn.append(f"[{name}] 取り込めた設問が 0 件")
        # ★0件でも捨てない。捨てると dropped（落とした理由）ごと消えて、
        #   「そのシートに何かあった」ことすら分からなくなる。
        #   SQL 側で「設問が1件も無い回は作らない」ので、空の回がDBにできる心配はない。
        sets.append(s)

    if only:
        missing = [n for n in only if n not in wb.sheetnames]
        if missing:
            warn.append(f"--sheet で指定したシートが無い: {missing} / このブックのシート = {wb.sheetnames}")

    # ★同じ課の範囲を指すシートが2枚あると、そのまま流すと二重に登録される。
    #   実データの「①1-3」と「①1-3 （新）」がこれ。どちらを使うかは人が決める。
    by_lesson = collections.defaultdict(list)
    for s in sets:
        if s["lesson"]:
            by_lesson[s["lesson"]].append(s["sheet"])
    for lesson, sheets in by_lesson.items():
        if len(sheets) > 1:
            warn.append(f"★同じ範囲「{lesson}課」のシートが {len(sheets)} 枚ある: {sheets}"
                        f"＝そのまま流すと二重に登録される。--sheet でどれか1枚を選ぶこと")

    # ★同じ問題の「別版」を見つける。実データでは 284 本のブックに
    #   「ルビあり」「ルビなし」（表記ゆれ＝ルビふり／ルビ振り）が並んでいて、
    #   中身は**同じ問題のルビ違い**だった（解答・配点・カテゴリは完全に一致）。
    #   シート名で決め打ちすると表記ゆれで漏れるので、**中身で判定する**。
    sig: dict[tuple, list[str]] = collections.defaultdict(list)
    for s in sets:
        if s["questions"]:
            key = tuple((q["seq"], q["correct_idx"], len(q["choices"])) for q in s["questions"])
            sig[key].append(s["sheet"])
    for sheets in sig.values():
        if len(sheets) > 1:
            warn.append(f"★同じ問題の別版とみられるシートが {len(sheets)} 枚: {sheets}"
                        f"（問題数・正解の並び・選択肢の数がすべて同じ）"
                        f"＝両方入れると二重になる。どちらを使うか決めて --sheet で選ぶこと")

    return sets, warn


# ---------------------------------------------------------------- SQL
def sq(s: str) -> str:
    return "'" + s.replace("'", "''") + "'"


def nq(v) -> str:
    """null 許容の値を SQL に。空・None は null。"""
    if v is None or v == "":
        return "null"
    return str(v) if isinstance(v, int) else sq(str(v))


def to_sql(sets: list[dict], answer_col: str, schema: str = SCHEMA_NEW) -> str:
    # ★設問が1件も無い回は作らない（空の回が学生の一覧に並んでしまうため）
    sets = [s for s in sets if s["questions"]]
    n = sum(len(s["questions"]) for s in sets)
    new = schema == SCHEMA_NEW
    L = [
        "-- 課題登録FMT の Excel → 学生ポータル（scripts/import_fmt_xlsx.py が生成）",
        f"-- 回 {len(sets)} 件 / 設問 {n} 件 / 正解列 = {answer_col} / スキーマ = {schema}",
        "-- ★ is_open = false で作る。中身を確かめてから教師画面で公開すること",
        "-- ★ 流す前に: 正解列の名前が実際のスキーマと合っているかだけ確認（--answer-col で変えられる）",
    ]
    if new:
        L += [
            "-- ★ db/2026-09-10_question_columns.sql を先に流しておくこと",
            "--    （image_name / category / points / explanation の4列を使う）",
            "-- ⚠ 解説は questions ではなく question_answers に入れる＝**学生から読めない表**。",
            "--    questions に置くと、公開中の回の解説を受験前に読めてしまう。",
            "-- ⓘ quiz_set_questions への登録は、db/2026-09-10_four_layers.sql の",
            "--    トリガ questions_sync_set_link が自動でやる（ここでは書かない）。",
        ]
    else:
        L.append("-- ⚠ --schema 2026-09-06: 画像・カテゴリ・配点・解説は**捨てられる**（列が無いため）")
    L += ["begin;", "do $$", "declare v_set uuid; v_q uuid;", "begin"]

    for s in sets:
        L.append(f"  -- ---- {s['sheet']} ----")
        L.append("  insert into quiz_sets(title, lesson, is_open) values "
                 f"({sq(s['title'])}, {sq(s['lesson'])}, false) returning id into v_set;")
        for it in s["questions"]:
            cols, vals = "quiz_set_id, seq, prompt", f"v_set, {it['seq']}, {sq(it['prompt'])}"
            if new:
                cols += ", image_name, category, points"
                vals += f", {nq(it['image_name'])}, {nq(it['category'])}, {nq(it['points'])}"
            L.append(f"  insert into questions({cols}) values ({vals}) returning id into v_q;")
            ch = ", ".join(f"(v_q, {i}, {sq(lab)})" for i, lab in enumerate(it["choices"], start=1))
            L.append(f"  insert into question_choices(question_id, idx, label) values {ch};")
            # ★正解は選択肢を入れたあと（外部キーが選択肢を指しているため）
            acols, avals = f"question_id, {answer_col}", f"v_q, {it['correct_idx']}"
            if new:
                acols += ", explanation"
                avals += f", {nq(it['explanation'])}"
            L.append(f"  insert into question_answers({acols}) values ({avals});")

    L += ["end $$;", "commit;", ""]
    return "\n".join(L)


# ---------------------------------------------------------------- main
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="課題登録FMT の Excel を学生ポータルの形へ変換する（DBには触らない）")
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--sheet", action="append", default=[],
                    help="取り込むシート名（複数可）。省略時はテンプレート以外の全シート")
    ap.add_argument("--exclude", action="append", default=[], help="除くシート名（複数可）")
    ap.add_argument("--include-template", action="store_true",
                    help="シート名に FMT を含むテンプレート本体も取り込む（既定は取り込まない）")
    ap.add_argument("--title-prefix", default="", help="quiz_sets.title の前に付ける文字")
    ap.add_argument("--answer-col", default=DEFAULT_ANSWER_COL)
    ap.add_argument("--ruby", choices=("keep", "strip", "html"), default="keep",
                    help="ルビ ${漢字}(よみ) の扱い。keep=そのまま（既定）／strip=消す／html=<ruby>タグ")
    ap.add_argument("--schema", choices=(SCHEMA_NEW, SCHEMA_OLD), default=SCHEMA_NEW,
                    help="出力するスキーマの世代。既定 2026-09-10＝画像・カテゴリ・配点・解説も書く"
                         "（db/2026-09-10_question_columns.sql を先に流しておくこと）。"
                         "2026-09-06＝それらの列が無い相手向け（★その4つは捨てられる）")
    ap.add_argument("--allow-duplicate-lesson", action="store_true",
                    help="同じ課の範囲のシートが複数あっても SQL を書く（既定は書かない）")
    ap.add_argument("--out-json", type=Path)
    ap.add_argument("--out-sql", type=Path)
    a = ap.parse_args(argv)

    if not a.xlsx.exists():
        print(f"ファイルが無い: {a.xlsx}", file=sys.stderr)
        return 2

    sets, warn = build(a.xlsx, a.sheet or None, a.exclude or None,
                       a.title_prefix, a.include_template, a.ruby)

    total = sum(len(s["questions"]) for s in sets)
    print(f"{a.xlsx.name}")
    print(f"取り込めた回 {len(sets)} 件 / 設問 {total} 件（is_open = false）\n")
    for s in sets:
        img = sum(1 for q in s["questions"] if q["image_name"])
        dist = collections.Counter(len(q["choices"]) for q in s["questions"])
        d = " / ".join(f"{k}択={v}問" for k, v in sorted(dist.items()))
        ng = len(s["dropped"])
        print(f"  ・{s['sheet']:<16} {len(s['questions']):3d}問  [{d}]"
              f"{f'  画像{img}問' if img else ''}"
              f"{f'  🔴取り込めず{ng}問' if ng else ''}"
              f"  → title={s['title']!r} lesson={s['lesson']!r}")

    dup_lesson = any(("同じ範囲" in w or "同じ問題の別版" in w) for w in warn)
    if warn:
        print(f"\n⚠ 目を通すこと {len(warn)} 件:")
        for w in warn:
            print("   " + w)
    else:
        print("\n検査: 問題なし")

    if a.out_json:
        a.out_json.parent.mkdir(parents=True, exist_ok=True)
        a.out_json.write_text(json.dumps(
            {"source": a.xlsx.name, "sets": sets}, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nJSON を書いた: {a.out_json}（画像の名前もここに残る）")

    if a.out_sql:
        if dup_lesson and not a.allow_duplicate_lesson:
            # ★二重登録は起きてからでは戻せないので、ここで止める
            print("\n🔴 SQL は書きませんでした＝同じ課の範囲のシートが複数あります。"
                  "\n   --sheet でどれか1枚を選ぶか、それでよければ --allow-duplicate-lesson を付けてください。")
            return 1
        a.out_sql.parent.mkdir(parents=True, exist_ok=True)
        a.out_sql.write_text(to_sql(sets, a.answer_col, a.schema), encoding="utf-8")
        print(f"SQL を書いた: {a.out_sql}（スキーマ {a.schema}）")
        if a.schema == SCHEMA_OLD:
            print("   ⚠ 画像・カテゴリ・配点・解説は SQL に出していません（この世代には列が無いため）")

    if not a.out_json and not a.out_sql:
        print("\n（--out-json / --out-sql を付けるとファイルに書きます）")
    return 1 if warn else 0


if __name__ == "__main__":
    sys.exit(main())
