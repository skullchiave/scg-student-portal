# -*- coding: utf-8 -*-
r"""make_sakumon_template.py — 配る「作問シート」のテンプレートを作る（2026-09-13）

★2026-09-13 きあ決定: 呼び名を「課題登録FMT」から **【作問シート】** に変えた。
  「課題登録」はヨリソルに登録するための名前で、2027年3月に意味を失う。
  「作問」は工程の名前なので、行き先が変わっても古びない。

  py -X utf8 scripts\make_sakumon_template.py
      … ドライブの配布先に 【作問シート】テンプレート（コピーして使う）.xlsx を書く

  py -X utf8 scripts\make_sakumon_template.py --out "どこか\ファイル.xlsx"

■ 列は「課題登録FMT」とまったく同じ（13列）
  中身は同じで、呼び方だけ変えた。既存331本のシートは触らない（きあ指示）。
  ★列を足したり順番を変えたりしないこと。取り込み口は **位置で読む**。

■ シート名に「コピーして」を入れてある
  取り込み口は、この語（と旧来の FMT）を **テンプレートらしさの手がかり** にする。
  🔴 ただし捨てるかどうかは **中身** で決める（見本と全部同じときだけ取り込まない）。
     ＝先生がうっかりこのシートに直接書いても、**捨てられない**。
       2026-09-13 に、名前だけで捨てていたせいで 120問 が落ちていたのを直した。

■ 見本の3行について
  この3行の指紋（FNV-1a）を scripts/import_fmt_xlsx.py と src/assets/fmt-import.js の
  TEMPLATE_SAMPLE_KEYS に入れてある。**入れ忘れると見本が本物として取り込まれる。**
  行を書き換えたら、必ず `--keys` で指紋を出し直して両方に貼ること。

  py -X utf8 scripts\make_sakumon_template.py --keys   … 指紋だけ出す（ファイルは書かない）
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "scripts"))

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    print("openpyxl が要ります:  py -m pip install openpyxl", file=sys.stderr)
    raise

import import_fmt_xlsx as fx  # noqa: E402

DEFAULT_OUT = Path(
    r"I:\マイドライブ\claude作業場\0.2 claude-work\【学生ポータル】 Student Portal"
    r"\【作問シート】テンプレート（コピーして使う）.xlsx")

SHEET_NAME = "作問シート（コピーして使う）"

HEADERS = ["問題番号", "問題文1", "問題文2", "添付ファイル名",
           "選択肢1", "選択肢2", "選択肢3", "選択肢4", "選択肢5",
           "解説", "解答", "カテゴリ", "配点"]

# 見本の3行。★2択・4択・ルビと解説、の3通りを見せる
SAMPLES = [
    # seq, 問題文1, 問題文2, 添付, 選択肢1..5, 解説, 解答, カテゴリ, 配点
    [1, "これは 見本です。この行を 消してから、問題を 書いてください。", "", "",
     "はい", "いいえ", "", "", "", "", 1, "文法", 1],
    [2, "つぎの ぶんの （  ）に 入る ことばは どれですか。",
     "わたしは まいにち コーヒー（  ）のみます。", "",
     "を", "が", "に", "で", "", "", 1, "文法", 1],
    [3, "${昨日}(きのう)、なにを たべましたか。", "", "",
     "たべました", "たべます", "", "", "",
     "「きのう」は すぎた ことなので、「〜ました」を つかいます。", 1, "文法", 1],
]


def sample_keys() -> list[str]:
    """見本3行の指紋。★取り込み口と同じ作り方（問題文1と2を改行でつなぐ）。"""
    keys = []
    for r in SAMPLES:
        t1, t2 = str(r[1]), str(r[2])
        prompt = "\n".join(x for x in (t1, t2) if x)
        choices = [str(c) for c in r[4:9] if str(c).strip()]
        keys.append(fx.question_key(prompt, choices))
    return keys


def build(out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    head_fill = PatternFill("solid", fgColor="1E40AF")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    sample_fill = PatternFill("solid", fgColor="FFF7E6")
    for r, row in enumerate(SAMPLES, start=2):
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.fill = sample_fill
            c.alignment = Alignment(vertical="top", wrap_text=(i in (2, 3, 10)))

    widths = [9, 38, 30, 16, 14, 14, 14, 14, 14, 32, 7, 12, 7]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--keys", action="store_true", help="見本の指紋だけ出す（ファイルは書かない）")
    a = ap.parse_args()

    keys = sample_keys()
    if a.keys:
        print("■ 見本3行の指紋（TEMPLATE_SAMPLE_KEYS に足すこと）")
        for k in keys:
            print(f'    "{k}",')
        missing = [k for k in keys if k not in fx.TEMPLATE_SAMPLE_KEYS]
        print(("🔴 まだ入っていない: " + ", ".join(missing)) if missing
              else "✅ 3つとも import_fmt_xlsx.py に入っています")
        return 1 if missing else 0

    build(a.out)
    print(f"✅ 書きました: {a.out}")
    print(f"   シート名: {SHEET_NAME}")
    print(f"   列: {len(HEADERS)}列（課題登録FMT と同じ位置）")
    print(f"   見本: {len(SAMPLES)}行")
    missing = [k for k in keys if k not in fx.TEMPLATE_SAMPLE_KEYS]
    if missing:
        print("\n🔴 見本の指紋が取り込み口に入っていません。このままだと **見本が本物として取り込まれます**。")
        print("   scripts/import_fmt_xlsx.py と src/assets/fmt-import.js の")
        print("   TEMPLATE_SAMPLE_KEYS に、次の値を足してください:")
        for k in missing:
            print(f'       "{k}",')
        return 1
    print("✅ 見本の指紋は取り込み口に入っています（このテンプレを取り込んでも0件になります）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
