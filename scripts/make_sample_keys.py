# -*- coding: utf-8 -*-
"""見本（白紙テンプレート）の設問の指紋を取り出す。
🔴 本文は印字しない。FNV-1a 32bit の16進だけ（本文には戻せない）。"""
import sys
import collections
sys.path.insert(0, r"I:\マイドライブ\claude作業場\0.4 repos\scg-student-portal\scripts")
import import_fmt_xlsx as fx
import audit_fmt as af
from openpyxl import load_workbook


def fnv1a(s: str) -> str:
    h = 0x811c9dc5
    for b in s.encode("utf-8"):
        h ^= b
        h = (h * 0x01000193) & 0xFFFFFFFF
    return format(h, "08x")


def qkey(q) -> str:
    """ルビ記法は外してから指紋にする（ルビあり版／なし版で同じ指紋になるように）。"""
    txt = fx.RUBY.sub(r"\1", q["prompt"] or "")
    chs = [fx.RUBY.sub(r"\1", c) for c in q["choices"]]
    return fnv1a(txt + "\u0001" + "\u0001".join(chs))


files = [p for p in af.find_xlsx(af.DEFAULT_BASE, af.SKIP_DIRS) if p.name not in af.EXCLUDE_NAMES]
sheets = []
for p in files:
    try:
        wb = load_workbook(p, data_only=True)
    except Exception:
        continue
    for name in wb.sheetnames:
        if fx.TEMPLATE_MARK not in name and "コピーして" not in name:
            continue
        try:
            s, _ = fx.build_sheet(wb[name], name, "", "keep")
        except Exception:
            continue
        if s.get("error") or not s.get("questions"):
            continue
        book = str(p.relative_to(af.DEFAULT_BASE)).split(chr(92))[0]
        sheets.append((name, frozenset(qkey(q) for q in s["questions"]), book))
    wb.close()

# 「同じ中身が何枚もある」ものが見本。1枚しかないものは、誰かが書いた本物
cnt = collections.Counter(k for _, k, _b in sheets)
books = collections.defaultdict(set)
for _n, k, b in sheets: books[k].add(b)
print(f"■ テンプレらしき名前のシート {len(sheets)} 枚 / 中身の種類 {len(cnt)}\n")
samples = set()
for keys, n in cnt.most_common():
    names = sorted(set(nm for nm, k, _b in sheets if k == keys))
    nb = len(books[keys])
    # ★見分ける条件は「複数ある」ではなく **複数の教材にまたがって同じ**。
    #   同じ教材の中だけで重複しているのは、テンプレではなく**本物の重複**（2026-09-13 に踏んだ）。
    is_sample = nb >= 2
    tag = "★見本（教材をまたいで同じ）" if is_sample else ("  本物（同じ教材の中の重複 "+str(n)+"枚）" if n>=2 else "  本物（1枚だけ）")
    print(f"   {n:3d}枚  教材{nb}種  {len(keys):3d}問  {tag}   {names[0][:34]}")
    if is_sample:
        samples |= set(keys)

print(f"\n■ 見本として覚える指紋: {len(samples)} 個")
print("TEMPLATE_SAMPLE_KEYS = {")
for i, k in enumerate(sorted(samples)):
    end = "\n" if i % 6 == 5 else ""
    print(f'    "{k}",', end=end if end else " ")
print("\n}")
