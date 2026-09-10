#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""test_audit_fmt.py — 教材フォルダ検査（scripts/audit_fmt.py）の回帰テスト（2026-09-10）

**DB不要・実物の教材フォルダ不要**で走る（CIや会社PCでも回せるように）。
一時ディレクトリに openpyxl でダミーの課題登録FMT Excelを数枚作り、それを `--dir` に見立てて検査する。

見ているのは次の4つ。
  ① 不備の分類（KINDS）が「直してもらうもの🔴」と「ただの記録ⓘ」を混ぜていないか
  ② 母数（分母）が出力（テキスト・JSON・HTML）に入っているか
  ③ `--dir` に存在しないパスを渡したとき、0件ではなくエラーとして分かるか
     （＝「見つからない」と「見つかったが0件」を混同しない）
  ④ 生成された HTML に問題文・選択肢の中身（学生情報が混ざりうる自由記述）が入らないか
     （このリポは public。レポートはファイル名・シート名・問題番号・不備の種類だけを出す設計）

★ダミーの学籍番号は 2604999 のように **末尾999から降順の帯** を使う（001側は実在の番号と衝突するため）。

  py -X utf8 -m unittest discover -s tests -p "test_audit_fmt.py" -v
"""
from __future__ import annotations

import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import audit_fmt as af  # noqa: E402

HEAD = ["問題番号", "問題文1", "問題文2", "添付ファイル名",
        "選択肢1", "選択肢2", "選択肢3", "選択肢4", "選択肢5",
        "解説", "解答", "カテゴリ", "配点"]


def qrow(no, t2="つぎの ぶんを えらんで ください", choices=("あ", "い", "う"), ans=1,
         t1="", image="", explain="", cat="文法", points=5):
    """FMT 1行ぶん（13列）。test_import_fmt_xlsx.py の qrow と同じ形。"""
    ch = list(choices) + [""] * (5 - len(choices))
    return [no, t1, t2, image, *ch, explain, ans, cat, points]


def make_book(path: Path, sheets: dict, head=None) -> Path:
    """sheets = {シート名: [行, 行, ...]} で Excel を作る。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(list(head or HEAD))
        for r in rows:
            ws.append(list(r))
    wb.save(path)
    return path


class Tmp(unittest.TestCase):
    """`self.base` を「教材フォルダ」に見立てる。ブックは group（1階層目）/ ファイル名 で置く。"""

    def setUp(self):
        self._d = tempfile.TemporaryDirectory()
        self.dir = Path(self._d.name)
        self.base = self.dir / "教材フォルダ"
        self.base.mkdir()

    def tearDown(self):
        self._d.cleanup()

    def add_book(self, group: str, filename: str, sheets: dict, head=None) -> Path:
        return make_book(self.base / group / filename, sheets, head)

    def scan(self, skip_dirs=()):
        return af.scan(self.base, list(skip_dirs))


# ------------------------------------------------------------------ ① 分類が混ざっていない
class KindsTest(unittest.TestCase):
    def test_need_flag_matches_the_red_marker(self):
        """🔴 で始まる分類だけが need=True。ⓘ・★ に need=True が紛れていないか（逆も）。"""
        for _key, label, need in af.KINDS:
            self.assertEqual(need, label.startswith("🔴"),
                              f"{label!r} の need フラグが印と食い違っている")

    def test_fixable_reasons_are_flagged_red(self):
        cases = [
            "選択肢が2個（2個以上ないと出題できない）",
            "解答が空か、数字でない",
            "問題文1も問題文2も空",
            "問題番号が重複: [1, 2]",
            "計算結果が入っていない（3行）",
            "選択肢が重複している（['い', 'い']）",
        ]
        for text in cases:
            label, need = af.kind_of(text)
            self.assertTrue(need, f"{text!r} は🔴（直してもらうもの）のはずが need=False だった")
            self.assertTrue(label.startswith("🔴"), f"{text!r} → {label!r}")

    def test_record_only_reasons_are_not_flagged_red(self):
        cases = [
            "まだ書かれていない問題が3問（問8〜10）",
            "画像つきの設問（a.png）",
            "2 列目から始まっている（1列目は「課番号」）",
            "5列目の見出しが「(空)」＝想定は「選択肢1」",
            "問題番号が数字でないので飛ばした行: 22行目「集計」",
            "同じ範囲「1-3課」のシートが 2 枚ある",
            "同じ問題の別版とみられるシートが 2 枚",
        ]
        for text in cases:
            label, need = af.kind_of(text)
            self.assertFalse(need, f"{text!r} は記録のはずが🔴扱いになっていた（{label!r}）")


# ------------------------------------------------------------------ ② 母数が出力に入っている
class DenominatorTest(Tmp):
    def test_total_rows_equals_questions_plus_dropped_plus_unwritten(self):
        blank = [9, "", "", "", "", "", "", "", "", "", "", "", 1]   # 番号だけ＝作りかけ
        self.add_book("001.グループA", "b.xlsx", {
            "①1-3": [qrow(1), qrow(2, choices=("あ",), ans=1), blank],
        })
        data = self.scan()
        T = data["total"]
        self.assertEqual(T["rows"], T["questions"] + T["dropped"] + T["unwritten"])
        self.assertEqual(T["questions"], 1)
        self.assertEqual(T["dropped"], 1)
        self.assertEqual(T["unwritten"], 1)

    def test_summary_text_prints_the_denominator(self):
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1)]})
        data = self.scan()
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            af.print_summary(data)
        out = buf.getvalue()
        self.assertIn("Excel にあった問題行", out)
        self.assertIn(f"{data['total']['rows']:6,}", out)

    def test_html_states_the_denominator_not_just_the_numerator(self):
        self.add_book("001.グループA", "b.xlsx", {
            "①1-3": [qrow(1), qrow(2, choices=("あ",), ans=1)],
        })
        data = self.scan()
        h = af.render_html(data)
        self.assertIn(f"{data['total']['rows']:,}", h)   # 分母
        self.assertIn(f"{data['total']['dropped']:,}", h)  # 分子だけでなく分母つきで出す
        self.assertIn("これが下の割合の母数", h)

    def test_json_carries_the_same_total(self):
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1)]})
        data = self.scan()
        rt = json.loads(json.dumps(data, ensure_ascii=False))
        self.assertEqual(rt["total"]["rows"], data["total"]["rows"])


# ------------------------------------------------------------------ ③ 「見つからない」と「0件」を混同しない
class MissingDirTest(Tmp):
    def run_main(self, argv):
        out_buf, err_buf = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out_buf), contextlib.redirect_stderr(err_buf):
            rc = af.main(argv)
        return rc, out_buf.getvalue(), err_buf.getvalue()

    def test_nonexistent_dir_is_an_error_not_zero_results(self):
        missing = self.base / "存在しないフォルダ"
        rc, _out, err = self.run_main(["--dir", str(missing), "--out", str(self.dir / "o.html")])
        self.assertEqual(rc, 2)
        self.assertIn("見つからない", err)
        self.assertFalse((self.dir / "o.html").exists())   # ★何も書かずに止まる

    def test_dir_that_is_actually_a_file_is_also_an_error(self):
        f = self.base / "not_a_dir.txt"
        f.write_text("x", encoding="utf-8")
        rc, _out, err = self.run_main(["--dir", str(f), "--out", str(self.dir / "o.html")])
        self.assertEqual(rc, 2)
        self.assertTrue(err.strip())

    def test_existing_but_empty_dir_is_not_an_error(self):
        """★フォルダはあるが xlsx が1本も無い＝これは正当な0件。エラーにしない。"""
        empty = self.base / "空フォルダ"
        empty.mkdir()
        rc, _out, _err = self.run_main(["--dir", str(empty), "--out", str(self.dir / "o.html")])
        self.assertEqual(rc, 0)
        self.assertTrue((self.dir / "o.html").exists())
        data = json.loads((self.dir / "o.json").read_text(encoding="utf-8"))
        self.assertEqual(data["total"]["books"], 0)


# ------------------------------------------------------------------ ④ HTML に設問の中身を出さない
class NoContentLeakTest(Tmp):
    def test_prompt_text_never_reaches_the_html(self):
        """問題文・選択肢の生の文字列はレポートに出さない設計を固定する（public リポ）。
        ダミーの学籍番号は 2604999（末尾999から降順の帯）で衝突を避ける。"""
        marker = "テスト太郎さん(学籍番号2604999)の話し合いより"
        self.add_book("001.グループA", "b.xlsx", {
            "①1-3": [qrow(1, t2=marker, choices=("あ", "い"), ans=1)],
        })
        data = self.scan()
        h = af.render_html(data)
        self.assertNotIn(marker, h)
        self.assertNotIn("2604999", h)
        # scan() の明細のどのフィールドにも生の問題文は乗らない設計であることも確認する
        self.assertNotIn(marker, json.dumps(data, ensure_ascii=False))

    def test_report_only_carries_structural_fields(self):
        """レポートに出るのは ファイル名・シート名・問題番号・不備の種類 だけ。"""
        self.add_book("001.グループA", "b.xlsx", {
            "①1-3": [qrow(1, choices=("あ",), ans=1)],   # 選択肢1個＝不備で落ちる
        })
        data = self.scan()
        h = af.render_html(data)
        self.assertIn("b.xlsx", h)
        self.assertIn("①1-3", h)
        self.assertIn("問1", h)


# ------------------------------------------------------------------ 走査そのもの
class ScanTest(Tmp):
    def test_skip_dirs_excludes_by_default(self):
        self.add_book(af.SKIP_DIRS[0], "旧テキスト.xlsx", {"①1-3": [qrow(1)]})
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1)]})
        data = self.scan(skip_dirs=af.SKIP_DIRS)
        self.assertEqual(data["total"]["books"], 1)

    def test_include_skipped_flag_reads_everything(self):
        """--include-skipped は main() 側で skip_dirs=[] を渡す形（スクリプト側で確認）。"""
        self.add_book(af.SKIP_DIRS[0], "旧テキスト.xlsx", {"①1-3": [qrow(1)]})
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1)]})
        data = self.scan(skip_dirs=[])
        self.assertEqual(data["total"]["books"], 2)

    def test_include_skipped_cli_flag(self):
        self.add_book(af.SKIP_DIRS[0], "旧テキスト.xlsx", {"①1-3": [qrow(1)]})
        out = self.dir / "o.html"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            af.main(["--dir", str(self.base), "--out", str(out), "--include-skipped"])
        data = json.loads(out.with_suffix(".json").read_text(encoding="utf-8"))
        self.assertEqual(data["total"]["books"], 1)

    def test_lock_file_and_progress_sheet_are_excluded(self):
        make_book(self.base / "001.グループA" / "~$b.xlsx", {"①1-3": [qrow(1)]})
        make_book(self.base / "000. コンテンツ移行進捗管理.xlsx", {"シート1": [qrow(1)]})
        self.add_book("001.グループA", "本物.xlsx", {"①1-3": [qrow(1)]})
        data = self.scan()
        self.assertEqual(data["total"]["books"], 1)
        self.assertIn("本物.xlsx", data["books"][0]["file"])
        self.assertIn("001.グループA", data["books"][0]["file"])

    def test_old_xls_is_counted_separately_not_silently_dropped(self):
        """openpyxl は .xls を読めないので検査対象には入らない。★黙って消えないよう別枠で数える。"""
        old = self.base / "001.グループA" / "古い形式.xls"
        old.parent.mkdir(parents=True, exist_ok=True)
        old.write_bytes(b"dummy")   # 中身は問わない（拡張子だけで別枠に振り分ける）
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1)]})
        data = self.scan()
        self.assertEqual(data["total"]["books"], 1)          # .xls は母数に混ざらない
        self.assertEqual(len(data["old_xls_files"]), 1)
        self.assertIn("古い形式.xls", data["old_xls_files"][0])

    def test_broken_book_is_recorded_in_failed_not_swallowed(self):
        broken = self.base / "001.グループA" / "壊れている.xlsx"
        broken.parent.mkdir(parents=True, exist_ok=True)
        broken.write_bytes(b"not a real xlsx file")
        data = self.scan()
        self.assertEqual(len(data["failed"]), 1)
        self.assertIn("壊れている.xlsx", data["failed"][0]["file"])

    def test_broken_workbook_warning_reaches_the_fix_list(self):
        """★import_fmt_xlsx.py の stale_cache() だけは文言の先頭に🔴が付く
        （他の警告は `[シート名] ...` から始まる）。シート名の抽出でここを弾くと、
        必須修正（need=True）なのに「直していただきたいもの」一覧から静かに漏れる。"""
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("①1-3")
        ws.append(list(HEAD))
        for i in range(2, 5):
            ws.append([f"={i - 1}", "", "もんだい", "", "あ", "い", "", "", "", "", 1, "文法", 1])
        p = self.base / "001.グループA" / "formula.xlsx"
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(p)

        data = self.scan()
        self.assertTrue(any("計算結果が入っていない" in w["text"] for w in data["warnings"]))
        h = af.render_html(data)
        self.assertIn("計算結果", h)   # ★一覧に出ていることを確認（漏れていたら消える）

    def test_kinds_counter_matches_red_vs_record_split(self):
        """パイプライン全体を通して、🔴 と ⓘ が kinds に別々に積まれるか。"""
        self.add_book("001.グループA", "b.xlsx", {
            "①1-3": [qrow(1, choices=("あ",), ans=1),      # 🔴 選択肢1個 → 不備
                      qrow(2, image="a.png")],              # ⓘ 画像つき → 記録のみ
        })
        data = self.scan()
        red = [k for k in data["kinds"] if k.startswith("🔴")]
        rec = [k for k in data["kinds"] if not k.startswith("🔴")]
        self.assertTrue(red)
        self.assertTrue(rec)


# ------------------------------------------------------------------ main（入口・1コマンド）
class MainTest(Tmp):
    def test_default_json_out_is_derived_from_out(self):
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1)]})
        out = self.dir / "レポート.html"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            af.main(["--dir", str(self.base), "--out", str(out)])
        self.assertTrue(out.exists())
        self.assertTrue(out.with_suffix(".json").exists())

    def test_explicit_json_out_is_honored(self):
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1)]})
        out = self.dir / "レポート.html"
        j = self.dir / "べつの場所.json"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            af.main(["--dir", str(self.base), "--out", str(out), "--json-out", str(j)])
        self.assertTrue(j.exists())

    def test_clean_scan_returns_0(self):
        self.add_book("001.グループA", "b.xlsx", {"①1-3": [qrow(1), qrow(2)]})
        out = self.dir / "o.html"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = af.main(["--dir", str(self.base), "--out", str(out)])
        self.assertEqual(rc, 0)

    def test_scan_with_defects_returns_nonzero(self):
        self.add_book("001.グループA", "b.xlsx", {
            "①1-3": [qrow(1, choices=("あ",), ans=1)],
        })
        out = self.dir / "o.html"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = af.main(["--dir", str(self.base), "--out", str(out)])
        self.assertEqual(rc, 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
